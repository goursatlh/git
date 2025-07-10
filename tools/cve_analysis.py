#!/usr/bin/env python3
import subprocess
import time
import sys
import os
import shutil
from datetime import datetime
import re
import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

current_time = datetime.now().strftime('%Y%m%d_%H%M%S')
folder_name = f'html_bak_{current_time}'

def gen_execl():
    # Get the current working directory
    #directory = os.getcwd()
    #directory = f'html_bak_20240821_164255'
    directory = folder_name 
    print("Html files directory:", directory)
    # Directory containing HTML files
    #directory = '/mnt/share/CVE'

    # Load the additional Excel file if provided
    if len(sys.argv) > 2:
        excel_file = sys.argv[2]
        print("Comparing execl file: ", excel_file)
        xx_df = pd.read_excel(excel_file)
    else:
        xx_df = pd.DataFrame()

    # Lists to store extracted data
    vuln_ids = []
    summaries = []
    components = []
    versions = []
    severities = []
    published_dates = []
    scan_dates = []
    analyses = []
    final_dispositions = []
    newly_updated_rows = []

    # Process each HTML file in the directory
    for filename in os.listdir(directory):
        if filename.endswith('.html'):
            filepath = os.path.join(directory, filename)
            with open(filepath, 'r', encoding='utf-8') as file:
                soup = BeautifulSoup(file, 'html.parser')

                # Extract all Vuln IDs and Summaries
                vuln_id_elements = soup.find_all(attrs={'data-testid': re.compile(r'vuln-detail-link-\d+')})
                summary_elements = soup.find_all(attrs={'data-testid': re.compile(r'vuln-summary-\d+')})

                # Extract component and version from filename
                fields = filename.split(' ')
                component = fields[2] if len(fields) >= 4 else ''
                version = fields[3] if len(fields) >= 4 else ''

                # Process each Vuln ID element
                for i, vuln_id_element in enumerate(vuln_id_elements):
                    vuln_id = vuln_id_element.text
                    vuln_ids.append(vuln_id)
                    summaries.append(summary_elements[i].text if i < len(summary_elements) else '')
                    
                    # Extract severity and keep only the text part behind the numbers
                    severity_element = soup.find(attrs={'data-testid': f'vuln-cvss3-link-{i}'})
                    if not severity_element:
                        severity_element = soup.find(attrs={'data-testid': f'vuln-cvss2-link-{i}'})
                    severity_text = severity_element.text if severity_element else ''
                    
                    # Remove numbers and keep only the text part behind it
                    severity_text_only = re.sub(r'^\d+(\.\d+)?\s*', '', severity_text)
                    
                    severities.append(severity_text_only)
                    
                    components.append(component)
                    versions.append(version)
                    
                    # Extract published date
                    published_date_element = soup.find(attrs={'data-testid': f'vuln-published-on-{i}'})
                    published_date = published_date_element.text if published_date_element else ''
                    
                    # Search for CVE ID in the additional Excel file
                    if not xx_df.empty:
                        match = xx_df[xx_df['CVE ID'] == vuln_id]
                        if not match.empty:
                            scan_date = match['Scan Date/Time'].values[0]
                            analysis = match['Analysis (including rationale why a CVE might not be applicable, ticket numbers for items fixed, comments on the use case or attack vector, etc.)'].values[0] if 'Analysis (including rationale why a CVE might not be applicable, ticket numbers for items fixed, comments on the use case or attack vector, etc.)' in match.columns else ''
                            final_disposition = match['Final Disposition'].values[0] if 'Final Disposition' in match.columns else ''
                            newly_updated_rows.append(False)
                        else:
                            scan_date = datetime.now().strftime('%Y-%m-%d %H:%M')
                            analysis = ''
                            final_disposition = ''
                            newly_updated_rows.append(True)
                    else:
                        scan_date = datetime.now().strftime('%Y-%m-%d %H:%M')
                        analysis = ''
                        final_disposition = ''
                        newly_updated_rows.append(False)
                    
                    published_dates.append(published_date)
                    print(scan_date)
                    scan_dates.append(scan_date)
                    analyses.append(analysis)
                    final_dispositions.append(final_disposition)

    # Create a DataFrame
    data = {
        'CVE ID': vuln_ids,
        'Component': components,
        'Component Version': versions,
        'Summary': summaries,
        'Date Published': published_dates,
        'Severity': severities,
        'Scan Date/Time': scan_dates,
        'Newly updated': newly_updated_rows,
        'Analysis (including rationale why a CVE might not be applicable, ticket numbers for items fixed, comments on the use case or attack vector, etc.)': analyses,
        'Final Disposition': final_dispositions
    }
    df = pd.DataFrame(data)

    # Sort the DataFrame by Component and CVE ID from small to large
    df.sort_values(by=['Component', 'CVE ID'], inplace=True)

    # Print DataFrame for debugging
    #print(df)

    # Generate a timestamped filename
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    excel_path = f'output_{timestamp}.xlsx'

    # Write the DataFrame to an Excel file
    df.to_excel(excel_path, index=False)

    # Load the workbook and select the active worksheet
    wb = load_workbook(excel_path)
    ws = wb.active

    # Define styles
    header_fill = PatternFill(start_color='70AD47', end_color='70AD47', fill_type='solid')  # Light green
    header_font = Font(bold=True, color='000000')  # Black font for header
    content_fills = [
        PatternFill(start_color='E8F3E1', end_color='E8F3E1', fill_type='solid'),  # Light greenish color
        PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid'),  # White color
    ]
    newly_updated_fill = PatternFill(start_color='FFFFD1', end_color='FFFFD1', fill_type='solid')  # Yellow color for newly updated rows
    alignment = Alignment(wrap_text=True, vertical='top')
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Apply styles to header
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = alignment
        cell.border = thin_border

    # Apply styles to content rows with alternating colors (E8F3E1 and white) and newly updated rows (yellow)
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column), start=2):
        fill = content_fills[(row_idx - 2) % len(content_fills)]

        # Check if the row is newly updated and apply the yellow fill to the entire row if true
        if row[-3].value:  # Assuming "Newly updated" is the third last column in the DataFrame/Excel sheet
            fill = newly_updated_fill

        for cell in row:
            cell.fill = fill
            cell.alignment = alignment
            cell.border = thin_border

    # Delete the "Newly updated" column from the worksheet before saving.
    ws.delete_cols(ws.max_column-2)

    # Adjust column width and apply text wrapping with a maximum width
    max_width = 100  # Set a maximum width for columns
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter  # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, max_width)
        ws.column_dimensions[column].width = adjusted_width
        for cell in col:
            cell.alignment = alignment

    # Save the workbook after deleting the "Newly updated" column.
    wb.save(excel_path)

    # Print statistics of newly added rows this time
    if len(sys.argv) > 2:
        newly_added_count = sum(newly_updated_rows)
    else:
        newly_added_count = len(newly_updated_rows)
    print()
    print(f"=====================================================================")
    print(f"Number of newly added rows this time: {newly_added_count}")

    print(f"Data has been written to {excel_path} with formatting.")
    print(f"=====================================================================")

def execute_command(para1, para2):
    command = f"python3 -c \"import nvd_cves; nvd_cves.get_batch('{para1} {para2}', verbose=1)\""
    subprocess.run(command, shell=True)

def move_html_files():
    if not os.path.exists(folder_name):
        os.makedirs(folder_name)
    
    for file in os.listdir('.'):
        if file.endswith('.html'):
            shutil.move(file, folder_name)

def main(filename):
    with open(filename, 'r') as file:
        for line in file:
            line = line.strip()
            if line:  # Filter out blank lines
                para1, para2 = line.split()
                execute_command(para1, para2)
                #time.sleep(2)  # Sleep for 2 seconds
    
    move_html_files()
    gen_execl()

if __name__ == "__main__":
    if len(sys.argv) > 3 or len(sys.argv) < 2 :
        print("Usage: ./cve_analysis.py <pair_filename> [<compare_execl_filename>]")
        sys.exit(1)
    
    filename = sys.argv[1]
    main(filename)
